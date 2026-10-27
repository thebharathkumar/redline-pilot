from dataclasses import dataclass, field

import docx
from docx.document import Document
from docx.oxml.ns import qn
from lxml import etree

REVISION_TAGS = ("w:ins", "w:del", "w:moveFrom", "w:moveTo")
TEXT_TAGS = {qn("w:t"): None, qn("w:tab"): "\t", qn("w:br"): "\n", qn("w:cr"): "\n"}


class ExistingRevisionsError(ValueError):
    pass


@dataclass
class RunSpan:
    element: etree._Element
    start: int
    end: int


@dataclass
class Paragraph:
    id: str
    element: etree._Element
    text: str
    runs: list[RunSpan] = field(default_factory=list)


@dataclass
class ParsedDocument:
    document: Document
    paragraphs: list[Paragraph]

    def __post_init__(self):
        self.by_id = {p.id: p for p in self.paragraphs}


def run_text(r: etree._Element) -> str:
    parts = []
    for child in r:
        if child.tag == qn("w:t"):
            parts.append(child.text or "")
        elif child.tag in TEXT_TAGS:
            parts.append(TEXT_TAGS[child.tag])
    return "".join(parts)


def owning_paragraph(el: etree._Element) -> etree._Element | None:
    parent = el.getparent()
    while parent is not None and parent.tag != qn("w:p"):
        parent = parent.getparent()
    return parent


def paragraph_runs(p: etree._Element) -> list[RunSpan]:
    spans = []
    offset = 0
    for r in p.iter(qn("w:r")):
        # Runs inside a text box belong to the text box's own paragraph.
        if owning_paragraph(r) is not p:
            continue
        text = run_text(r)
        spans.append(RunSpan(r, offset, offset + len(text)))
        offset += len(text)
    return spans


def in_text_box(p: etree._Element) -> bool:
    parent = p.getparent()
    while parent is not None:
        if parent.tag == qn("w:txbxContent"):
            return True
        parent = parent.getparent()
    return False


def find_revisions(body: etree._Element) -> list[str]:
    found = []
    for tag in REVISION_TAGS:
        if body.find(".//" + qn(tag)) is not None:
            found.append(tag)
    return found


def parse_document(document: Document) -> ParsedDocument:
    body = document.element.body
    existing = find_revisions(body)
    if existing:
        raise ExistingRevisionsError(
            "input already contains tracked changes (" + ", ".join(existing) + "); "
            "accept or reject them in Word and save a clean copy first"
        )
    paragraphs = []
    for p in body.iter(qn("w:p")):
        if in_text_box(p):
            continue
        runs = paragraph_runs(p)
        text = "".join(run_text(s.element) for s in runs)
        paragraphs.append(Paragraph(f"p{len(paragraphs) + 1:04d}", p, text, runs))
    return ParsedDocument(document, paragraphs)


def load(path: str) -> ParsedDocument:
    return parse_document(docx.Document(path))
