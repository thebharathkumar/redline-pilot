import docx
import pytest


def make_document():
    document = docx.Document()
    document.add_paragraph("Mutual Non-Disclosure Agreement")
    p = document.add_paragraph("The Recipient shall keep ")
    p.add_run("all Confidential Information").bold = True
    p.add_run(" strictly ")
    p.add_run("confidential and secure").italic = True
    p.add_run(" for five years.")
    document.add_paragraph("")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Governing law"
    table.cell(0, 1).text = "State of New York"
    document.add_paragraph("Signed by the parties.")
    return document


@pytest.fixture
def document():
    return make_document()
