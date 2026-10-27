"""Build fixtures/common_paper_nda.docx and fixtures/counterparty_nda.docx.

The counterparty variant is the base text with the deviations listed below applied,
saved as a clean document (no tracked changes). fixtures/DEVIATIONS.md is written
from the same list.
"""

import re
import sys
from pathlib import Path

import docx

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from redline.review import save_stable  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "fixtures/source/common_paper_mnda_v1.md"

INTRO = (
    "This Mutual Non-Disclosure Agreement (the “**MNDA**”) consists of: (1) this Cover Page "
    "(“**Cover Page**”) and (2) the Standard Terms below, which are based on the Common Paper "
    "Mutual NDA Standard Terms Version 1.0 (“**Standard Terms**”)."
)

COVER = [
    ("Purpose", "Evaluating whether to enter into a business relationship with the other party."),
    ("Effective Date", "January 5, 2026"),
    ("MNDA Term", "Expires 1 year from Effective Date."),
    (
        "Term of Confidentiality",
        "1 year from Effective Date, but in the case of trade secrets until Confidential Information "
        "is no longer considered a trade secret under applicable laws.",
    ),
    ("Governing Law", "New York"),
    ("Jurisdiction", "Courts located in New York County, New York."),
]

SIGNATURE_ROWS = ["Signature", "Print Name", "Title", "Company", "Notice Address", "Date"]

ATTRIBUTION = (
    "Standard Terms based on the Common Paper Mutual Non-Disclosure Agreement Version 1.0, "
    "free to use under CC BY 4.0."
)

# (rule_id, where, [(find, replace), ...]). Each find must occur exactly once in the base text.
DEVIATIONS = [
    (
        "ci-marking",
        "Section 1",
        [
            (
                "which (1) the Disclosing Party identifies to the receiving party (“**Receiving Party**”) as "
                "“confidential”, “proprietary”, or the like or (2) should be reasonably understood as "
                "confidential or proprietary due to its nature and the circumstances of its disclosure",
                "which is marked “confidential” or “proprietary” at the time of disclosure to the receiving "
                "party (“**Receiving Party**”) and, if disclosed orally, is summarized in writing and marked "
                "as confidential within ten (10) days of disclosure",
            ),
        ],
    ),
    (
        "ci-scope",
        "Section 1",
        [
            (
                "Confidential Information includes technical or business information, product designs or "
                "roadmaps, requirements, pricing, security and compliance documentation, technology, "
                "inventions and know-how.",
                "Confidential Information is limited to technical information, source code and trade secrets.",
            ),
        ],
    ),
    (
        "disc-advisors",
        "Section 2(b)",
        [
            (
                "to its employees, agents, advisors, contractors and other representatives having",
                "to its employees having",
            ),
        ],
    ),
    (
        "disc-legal",
        "Section 4",
        [
            (
                "reasonably cooperates, at the Disclosing Party’s expense, with the Disclosing Party’s efforts",
                "uses best efforts, at the Receiving Party’s expense, to contest the requirement and to support "
                "the Disclosing Party’s efforts",
            ),
        ],
    ),
    (
        "term-length",
        "Cover Page, Term of Confidentiality",
        [
            (
                "1 year from Effective Date, but in the case of trade secrets until Confidential Information "
                "is no longer considered a trade secret under applicable laws.",
                "In perpetuity.",
            ),
        ],
    ),
    (
        "ret-retention",
        "Section 6",
        [
            (
                "As an exception to subsection (b), the Receiving Party may retain Confidential Information in "
                "accordance with its standard backup or record retention policies or as required by law, but "
                "the terms of this MNDA will continue to apply to the retained Confidential Information.",
                "The Receiving Party will permanently delete all Confidential Information, including copies "
                "held in electronic backup systems.",
            ),
        ],
    ),
    (
        "rem-equitable",
        "Section 10",
        [
            (
                "the Disclosing Party is entitled to seek appropriate equitable relief, including an "
                "injunction, in addition to its other remedies.",
                "the Disclosing Party is entitled to an injunction without posting a bond, and the Receiving "
                "Party will pay liquidated damages of $50,000 for each breach, in addition to the Disclosing "
                "Party’s other remedies.",
            ),
        ],
    ),
    (
        "law-venue",
        "Cover Page, Governing Law and Jurisdiction",
        [
            (
                "Courts located in New York County, New York.",
                "Courts located in Travis County, Texas.",
            ),
            (
                "Governing Law | New York",
                "Governing Law | Texas",
            ),
        ],
    ),
]


def standard_terms() -> list[str]:
    text = SOURCE.read_text(encoding="utf-8")
    return [line for line in text.splitlines() if re.match(r"^\d+\. \*\*", line)]


def add_marked_text(paragraph, text: str) -> None:
    """Add text where **x** becomes a bold run."""
    for i, piece in enumerate(re.split(r"\*\*", text)):
        if piece:
            paragraph.add_run(piece).bold = i % 2 == 1


def flatten(cover, terms) -> str:
    return "\n".join([INTRO] + [f"{k} | {v}" for k, v in cover] + terms)


def unflatten(flat: str):
    lines = flat.split("\n")
    n = len(COVER)
    intro, cover, terms = lines[0], [tuple(l.split(" | ", 1)) for l in lines[1:1 + n]], lines[1 + n:]
    return intro, cover, terms


def build(intro, cover, terms, path: Path) -> None:
    document = docx.Document()
    document.core_properties.author = ""
    document.core_properties.title = "Mutual Non-Disclosure Agreement"
    document.add_paragraph("Mutual Non-Disclosure Agreement", style="Title")
    document.add_heading("Cover Page", level=1)
    add_marked_text(document.add_paragraph(), intro)
    table = document.add_table(rows=len(cover), cols=2)
    table.style = "Table Grid"
    for row, (label, value) in zip(table.rows, cover):
        row.cells[0].paragraphs[0].add_run(label).bold = True
        row.cells[1].paragraphs[0].add_run(value)
    document.add_paragraph("By signing this Cover Page, each party agrees to enter into this MNDA as of the Effective Date.")
    sig = document.add_table(rows=len(SIGNATURE_ROWS) + 1, cols=3)
    sig.style = "Table Grid"
    sig.cell(0, 1).text = "PARTY 1"
    sig.cell(0, 2).text = "PARTY 2"
    for i, label in enumerate(SIGNATURE_ROWS, start=1):
        sig.cell(i, 0).text = label
    document.add_heading("Standard Terms", level=1)
    for line in terms:
        add_marked_text(document.add_paragraph(), line)
    document.add_paragraph().add_run(ATTRIBUTION).italic = True
    save_stable(document, str(path))


def apply_deviations(flat: str) -> str:
    for rule_id, where, pairs in DEVIATIONS:
        for find, replace in pairs:
            count = flat.count(find)
            if count != 1:
                raise SystemExit(f"deviation {rule_id} ({where}): find occurs {count} times")
            flat = flat.replace(find, replace)
    return flat


def plain(text: str) -> str:
    return text.replace("**", "")


def write_deviations(path: Path) -> None:
    lines = [
        "# Counterparty deviations",
        "",
        "fixtures/counterparty_nda.docx is fixtures/common_paper_nda.docx with these changes, written by",
        "scripts/make_fixtures.py. Each row names the playbook rule it should trigger.",
        "",
    ]
    for n, (rule_id, where, pairs) in enumerate(DEVIATIONS, start=1):
        lines += [f"## {n}. {rule_id} ({where})", ""]
        for find, replace in pairs:
            lines += [f"Base: {plain(find)}", "", f"Counterparty: {plain(replace)}", ""]
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    base = flatten(COVER, standard_terms())
    build(*unflatten(base), ROOT / "fixtures/common_paper_nda.docx")
    build(*unflatten(apply_deviations(base)), ROOT / "fixtures/counterparty_nda.docx")
    write_deviations(ROOT / "fixtures/DEVIATIONS.md")


if __name__ == "__main__":
    main()
