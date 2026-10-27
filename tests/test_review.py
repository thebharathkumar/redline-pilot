import json
import shutil
import subprocess
from pathlib import Path

import docx
import pytest
from docx.oxml.ns import qn

from redline.cli import main
from redline.parse import load, parse_document
from redline.playbook import load_playbook
from redline.revisions import document_texts
from redline.schema import Edit, ReviewResult
from redline.validate import validate

ROOT = Path(__file__).resolve().parent.parent
CONTRACT = ROOT / "fixtures/counterparty_nda.docx"
PLAYBOOK = ROOT / "playbooks/nda.yaml"
REPLAY = ROOT / "tests/data/replay_response.json"
DATE = "2026-01-01T00:00:00Z"


def review(out: Path, response: Path = REPLAY) -> Path:
    code = main([
        "review", str(CONTRACT), "--playbook", str(PLAYBOOK), "--out", str(out),
        "--offline", str(response), "--date", DATE,
    ])
    assert code == 0
    return out


def check(paragraph_id, find, replace="x", rule_id="disc-legal"):
    document = load(str(CONTRACT))
    edit = Edit(paragraph_id=paragraph_id, rule_id=rule_id, find=find, replace=replace, comment="c")
    accepted, _, rejected = validate(ReviewResult(edits=[edit]), document, load_playbook(str(PLAYBOOK)))
    return accepted, [r.reason for r in rejected]


def test_find_found_once_is_accepted():
    accepted, reasons = check("p0042", "subpoena or court order")
    assert len(accepted) == 1 and reasons == []


def test_find_not_found_is_rejected():
    assert check("p0042", "arbitration panel")[1] == ["find not in paragraph"]


def test_find_found_twice_is_rejected():
    assert check("p0042", "Disclosing Party")[1] == ["find occurs 2 times"]


def test_find_must_be_whole_words():
    assert check("p0042", "isclose")[1] == ["find not in paragraph"]


def test_unknown_rule_and_paragraph_are_rejected():
    assert check("p0042", "subpoena", rule_id="nope")[1] == ["unknown rule_id"]
    assert check("p9999", "subpoena")[1] == ["unknown paragraph_id"]


def test_overlapping_edits_keep_the_first():
    document = load(str(CONTRACT))
    edits = [
        Edit(paragraph_id="p0042", rule_id="disc-legal", find="subpoena or court order", replace="a", comment="c"),
        Edit(paragraph_id="p0042", rule_id="disc-legal", find="court order, provided", replace="b", comment="c"),
    ]
    accepted, _, rejected = validate(ReviewResult(edits=edits), document, load_playbook(str(PLAYBOOK)))
    assert [a.edit.replace for a in accepted] == ["a"]
    assert rejected[0].reason == "overlaps disc-legal"


def test_review_round_trip(tmp_path):
    original = document_texts(docx.Document(str(CONTRACT)).element.body, accept=True)
    out = docx.Document(str(review(tmp_path / "out.docx")))
    body = out.element.body

    response = json.loads(REPLAY.read_text())
    expected = list(original)
    ids = [p.id for p in parse_document(docx.Document(str(CONTRACT))).paragraphs]
    kept = [e for e in response["edits"] if "should be dropped" not in e["comment"]]
    for e in kept:
        i = ids.index(e["paragraph_id"])
        expected[i] = expected[i].replace(e["find"], e["replace"])

    assert document_texts(body, accept=True) == expected
    assert document_texts(body, accept=False) == original

    marks = list(body.iter(qn("w:ins"), qn("w:del")))
    assert len({m.get(qn("w:id")) for m in marks}) == len(marks)
    assert all(d.find(".//" + qn("w:t")) is None for d in body.iter(qn("w:del")))

    comments = list(out.comments)
    assert len(comments) == len(kept) + 1
    assert all(c.text.startswith("[") for c in comments)
    assert comments[-1].text.startswith("[ret-retention] ")
    for c in comments:
        cid = str(c.comment_id)
        for tag in ("commentRangeStart", "commentRangeEnd", "commentReference"):
            assert len(body.xpath(f'.//w:{tag}[@w:id="{cid}"]')) == 1


def test_offline_review_is_byte_stable(tmp_path):
    first = review(tmp_path / "a.docx").read_bytes()
    second = review(tmp_path / "b.docx").read_bytes()
    assert first == second


@pytest.mark.skipif(not (ROOT / "fixtures/demo_response.json").exists(), reason="fixtures/demo_response.json not created yet")
def test_make_demo_is_byte_stable(tmp_path):
    outputs = []
    for _ in range(2):
        subprocess.run(["make", "-s", "demo"], cwd=ROOT, check=True, capture_output=True)
        outputs.append((ROOT / "out/demo_redlined.docx").read_bytes())
    assert outputs[0] == outputs[1]


def test_existing_revisions_exit_with_error(tmp_path, capsys):
    reviewed = review(tmp_path / "once.docx")
    code = main([
        "review", str(reviewed), "--playbook", str(PLAYBOOK), "--out", str(tmp_path / "twice.docx"),
        "--offline", str(REPLAY), "--date", DATE,
    ])
    assert code == 1
    assert "already contains tracked changes" in capsys.readouterr().err


@pytest.mark.skipif(shutil.which("soffice") is None, reason="soffice not installed")
def test_soffice_converts_output_to_pdf(tmp_path):
    out = review(tmp_path / "out.docx")
    subprocess.run(
        ["soffice", "--headless", f"-env:UserInstallation=file://{tmp_path}/lo", "--convert-to", "pdf",
         "--outdir", str(tmp_path), str(out)],
        check=True, capture_output=True, timeout=180,
    )
    pdf = tmp_path / "out.pdf"
    assert pdf.exists() and pdf.read_bytes().startswith(b"%PDF")
