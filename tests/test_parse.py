import pytest
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from redline.parse import ExistingRevisionsError, parse_document


def test_ids_are_sequential_and_include_table_cells(document):
    parsed = parse_document(document)
    assert [p.id for p in parsed.paragraphs] == ["p0001", "p0002", "p0003", "p0004", "p0005", "p0006"]
    assert [p.text for p in parsed.paragraphs][3:] == ["Governing law", "State of New York", "Signed by the parties."]
    assert parsed.by_id["p0005"].text == "State of New York"


def test_run_offsets_cover_paragraph_text(document):
    para = parse_document(document).by_id["p0002"]
    assert para.text == "The Recipient shall keep all Confidential Information strictly confidential and secure for five years."
    assert para.runs[0].start == 0
    for prev, cur in zip(para.runs, para.runs[1:]):
        assert prev.end == cur.start
    assert para.runs[-1].end == len(para.text)
    bold = para.runs[1]
    assert para.text[bold.start:bold.end] == "all Confidential Information"


def test_ids_are_stable_across_parses(document):
    first = [(p.id, p.text) for p in parse_document(document).paragraphs]
    second = [(p.id, p.text) for p in parse_document(document).paragraphs]
    assert first == second


@pytest.mark.parametrize("tag", ["w:ins", "w:del", "w:moveFrom", "w:moveTo"])
def test_existing_revisions_are_rejected(document, tag):
    p = document.paragraphs[0]._p
    p.append(OxmlElement(tag, attrs={qn("w:id"): "1", qn("w:author"): "x"}))
    with pytest.raises(ExistingRevisionsError, match=tag):
        parse_document(document)
